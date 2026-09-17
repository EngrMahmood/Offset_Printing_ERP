"""Excel import/export for the BOM app, matching the planning team's existing
Excel workflow (see the reference carton BOM export this format is modelled
on: Sku-FG / Bom Item Sku-RM / Consumption-RM / Wastage / Manufacturing Step /
Bom Item Type). The carton workbook itself is never imported — it's reference
only; these templates are for the offset department.
"""
from __future__ import annotations

import io
from decimal import Decimal, InvalidOperation

from django.http import HttpResponse

try:
    import openpyxl
    from openpyxl.styles import Font, PatternFill
    from openpyxl.worksheet.datavalidation import DataValidation
    EXCEL_AVAILABLE = True
except ImportError:
    EXCEL_AVAILABLE = False

RAW_ITEM_HEADERS = [
    'Item Code', 'Name', 'Category', 'Item Role', 'UOM', 'Specification',
    'Brand/Grade', 'GSM', 'Sheet Size', 'Unit Cost', 'Vendor',
    'Pack Size', 'MOQ', 'Safety Stock', 'Max Stock', 'Lead Time Days',
]

SKU_BOM_HEADERS = [
    'Sku-FG', 'Unit', 'Type', 'Production Line', 'Bom Status',
    'Bom Item Sku-RM', 'Bom Item Unit-RM', 'Consumption-RM', 'Wastage',
    'Manufacturing Step', 'Bom Item Type',
]


def _style_header(ws, ncols):
    fill = PatternFill(start_color='DDEBF7', end_color='DDEBF7', fill_type='solid')
    for col in range(1, ncols + 1):
        cell = ws.cell(row=1, column=col)
        cell.font = Font(bold=True)
        cell.fill = fill
    ws.freeze_panes = 'A2'


def _write_workbook(filename, headers, data_rows, example_row=None):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = 'Sheet1'
    ws.append(headers)
    if example_row:
        ws.append(example_row)
    for row in data_rows:
        ws.append(row)
    _style_header(ws, len(headers))
    for col in range(1, len(headers) + 1):
        ws.column_dimensions[openpyxl.utils.get_column_letter(col)].width = 22

    buffer = io.BytesIO()
    wb.save(buffer)
    buffer.seek(0)
    response = HttpResponse(
        buffer.read(),
        content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
    )
    response['Content-Disposition'] = f'attachment; filename="{filename}"'
    return response


def download_raw_items_template():
    example = ['INK-0001', 'Water Base Ink CL Black', 'Ink', 'ink', 'KG', '', 'CL', '', '', '0', '', '1', '0', '0', '0', '1']
    return _write_workbook('raw_items_template.xlsx', RAW_ITEM_HEADERS, [], example_row=example)


def download_sku_bom_template():
    example = [
        'CTN-001', 'PCS', 'OFFSET_BOX', 'OFFSET', 'DRAFT',
        'PAP-ART300-25X36', 'SHEET', '0.001', '0.05', 'PRINTING', 'Paper/Board',
    ]
    return _write_workbook('sku_bom_template.xlsx', SKU_BOM_HEADERS, [], example_row=example)


def export_raw_items(queryset, filename='raw_items.xlsx'):
    rows = []
    for item in queryset.select_related('category', 'uom', 'default_vendor'):
        rows.append([
            item.item_code, item.name, item.category.name, item.item_role, item.uom.code,
            item.specification, item.brand_grade, item.attributes.get('gsm', ''),
            item.attributes.get('sheet_size', ''), str(item.unit_cost),
            item.default_vendor.name if item.default_vendor_id else '',
            str(item.pack_size), str(item.moq), str(item.safety_stock),
            str(item.max_stock_level), item.lead_time_days,
        ])
    return _write_workbook(filename, RAW_ITEM_HEADERS, rows)


def export_sku_bom(sku_bom_queryset, filename='sku_boms.xlsx'):
    rows = []
    for bom in sku_bom_queryset.select_related('sku_recipe').prefetch_related('lines__raw_item', 'lines__process_step', 'lines__raw_item__category'):
        for line in bom.lines.all():
            rows.append([
                bom.sku_recipe.sku, 'PCS', bom.sku_recipe.product_type or '', 'OFFSET', bom.get_status_display().upper(),
                line.raw_item.item_code, line.uom_code, str(line.qty_per_unit), str(line.wastage_percent),
                line.process_step.code, line.raw_item.category.name,
            ])
    return _write_workbook(filename, SKU_BOM_HEADERS, rows)


def _row_to_dict(headers, row):
    return {headers[i]: row[i] for i in range(len(headers)) if i < len(row)}


def _read_rows(upload_file, expected_headers):
    wb = openpyxl.load_workbook(upload_file, data_only=True)
    ws = wb.active
    header_row = [str(c.value).strip() if c.value is not None else '' for c in ws[1]]
    rows = []
    for r in range(2, ws.max_row + 1):
        values = [ws.cell(r, c).value for c in range(1, len(header_row) + 1)]
        if not any(v is not None and str(v).strip() for v in values):
            continue
        rows.append(dict(zip(header_row, values)))
    return rows


def _dec(value, default='0'):
    if value is None or str(value).strip() == '':
        return Decimal(default)
    try:
        return Decimal(str(value))
    except InvalidOperation:
        return Decimal(default)


def preview_import_raw_items(upload_file):
    """Dry-run: returns (rows_preview, errors) without writing anything."""
    from .models import ItemCategory, RawItem, UnitOfMeasure

    rows = _read_rows(upload_file, RAW_ITEM_HEADERS)
    preview, errors = [], []
    for idx, row in enumerate(rows, start=2):
        item_code = str(row.get('Item Code') or '').strip()
        name = str(row.get('Name') or '').strip()
        category_name = str(row.get('Category') or '').strip()
        uom_code = str(row.get('UOM') or '').strip().upper()

        row_errors = []
        if not item_code:
            row_errors.append('Item Code is required')
        if not name:
            row_errors.append('Name is required')
        category = ItemCategory.objects.filter(name__iexact=category_name).first() if category_name else None
        if category_name and category is None:
            row_errors.append(f'Unknown category "{category_name}"')
        uom = UnitOfMeasure.objects.filter(code__iexact=uom_code).first() if uom_code else None
        if uom_code and uom is None:
            row_errors.append(f'Unknown UOM "{uom_code}"')

        exists = RawItem.objects.filter(item_code=item_code).exists() if item_code else False
        preview.append({
            'row': idx, 'item_code': item_code, 'name': name, 'action': 'update' if exists else 'create',
            'errors': row_errors, 'raw_row': row,
        })
        errors.extend(f'Row {idx}: {e}' for e in row_errors)
    return preview, errors


def commit_import_raw_items(upload_file, user=None):
    from django.db import transaction

    from .models import ItemCategory, RawItem, UnitOfMeasure

    preview, errors = preview_import_raw_items(upload_file)
    if errors:
        return None, errors

    created, updated = 0, 0
    with transaction.atomic():
        for entry in preview:
            row = entry['raw_row']
            category = ItemCategory.objects.filter(name__iexact=str(row.get('Category') or '').strip()).first()
            uom = UnitOfMeasure.objects.filter(code__iexact=str(row.get('UOM') or '').strip()).first()
            gsm = row.get('GSM')
            sheet_size = str(row.get('Sheet Size') or '').strip()
            attributes = {}
            if gsm not in (None, ''):
                try:
                    attributes['gsm'] = int(float(gsm))
                except (TypeError, ValueError):
                    pass
            if sheet_size:
                attributes['sheet_size'] = sheet_size

            defaults = {
                'name': entry['name'],
                'category': category,
                'uom': uom,
                'specification': str(row.get('Specification') or '').strip(),
                'brand_grade': str(row.get('Brand/Grade') or '').strip(),
                'attributes': attributes,
                'unit_cost': _dec(row.get('Unit Cost')),
                'pack_size': _dec(row.get('Pack Size'), '1'),
                'moq': _dec(row.get('MOQ')),
                'safety_stock': _dec(row.get('Safety Stock')),
                'max_stock_level': _dec(row.get('Max Stock')),
                'lead_time_days': int(_dec(row.get('Lead Time Days'), '1')),
                'created_by': user,
            }
            _, was_created = RawItem.objects.update_or_create(item_code=entry['item_code'], defaults=defaults)
            created += 1 if was_created else 0
            updated += 0 if was_created else 1

    return {'created': created, 'updated': updated}, []
